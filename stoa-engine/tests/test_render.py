"""Renderer smoke tests: files exist, xlsx has 3 sheets, each PDF > 5 pages,
watermark string present when sample_data is true, manifest hashes outputs."""

import json

from openpyxl import load_workbook
from pypdf import PdfReader

from stoa_engine import pipeline


def test_export_all(tmp_path, sample_path):
    result = pipeline.export(str(sample_path), "all", str(tmp_path))

    posture = tmp_path / "packet_posture.pdf"
    performance = tmp_path / "packet_performance.pdf"
    xlsx = tmp_path / "performance_data.xlsx"
    manifest = tmp_path / "manifest.json"
    submission_copy = tmp_path / "submission.json"

    for f in (posture, performance, xlsx, manifest, submission_copy):
        assert f.exists(), f"missing {f.name}"

    # each PDF > 5 pages
    assert len(PdfReader(str(posture)).pages) > 5
    assert len(PdfReader(str(performance)).pages) > 5

    # watermark present (sample data)
    txt = "\n".join((p.extract_text() or "") for p in PdfReader(str(posture)).pages)
    assert "SAMPLE — FICTIONAL DATA — NOT FOR SUBMISSION" in txt

    # xlsx has exactly the three required sheets
    wb = load_workbook(str(xlsx))
    assert wb.sheetnames == ["Performance Data", "Covered Models", "Aggregates"]
    # 2 banner rows + 1 header + 70 runs = 73
    assert wb["Performance Data"].max_row == 73
    assert wb["Performance Data"].cell(1, 1).value.startswith("SAMPLE")

    # manifest lists every output file with a sha256
    m = json.loads(manifest.read_text())
    files = {e["file"] for e in m["files"]}
    assert {"packet_posture.pdf", "packet_performance.pdf",
            "performance_data.xlsx", "submission.json"} <= files
    assert all(len(e["sha256"]) == 64 for e in m["files"])
    assert m["sample_data"] is True

    # gaps surfaced
    assert result["gaps"]


def test_export_single_template(tmp_path, sample_path):
    pipeline.export(str(sample_path), "posture", str(tmp_path))
    assert (tmp_path / "packet_posture.pdf").exists()
    # posture-only export renders no performance workbook
    assert not (tmp_path / "performance_data.xlsx").exists()
