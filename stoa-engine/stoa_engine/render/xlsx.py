"""XLSX renderer (openpyxl).

Three sheets:
  1. "Performance Data"  — one row per performance run, columns:
     Timestamp | System ID | Use Case ID | Predicted Value | Ground Truth Value
     | Financial Loss Proxy (US$). Header note flags SIMULATED evidence.
  2. "Covered Models"    — one row per system from covered_model_draft.
  3. "Aggregates"        — per-system aggregate metrics.

When meta.sample_data is true, every sheet carries the fictional-data
watermark in row 1. loss_proxy is always labeled a proxy, never actual loss.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

WATERMARK = "SAMPLE — FICTIONAL DATA — NOT FOR SUBMISSION"

_HEADER_FILL = PatternFill("solid", fgColor="182A3E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_NOTE_FONT = Font(color="B98A2F", bold=True, italic=True)
_WM_FONT = Font(color="B23A3A", bold=True)


def _banner(ws, text: str, span: int, font: Font) -> int:
    """Write a merged banner row; return the next free row index (1-based)."""
    ws.append([text])
    r = ws.max_row
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=span)
    ws.cell(row=r, column=1).font = font
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="left")
    return r + 1


def _write_header(ws, headers: list[str]) -> None:
    ws.append(headers)
    r = ws.max_row
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _autosize(ws, max_width: int = 60) -> None:
    for col_cells in ws.columns:
        # find the column letter from the first non-merged cell
        letter = None
        length = 0
        for cell in col_cells:
            if isinstance(cell.column, int):
                letter = get_column_letter(cell.column)
            val = cell.value
            if val is not None:
                length = max(length, len(str(val)))
        if letter:
            ws.column_dimensions[letter].width = min(max_width, max(12, length + 2))


def render_xlsx(submission: dict[str, Any], out_path: Path) -> Path:
    sample = submission.get("meta", {}).get("sample_data", False)
    systems = submission.get("systems", []) or []

    wb = Workbook()

    # ---- Sheet 1: Performance Data ----------------------------------------
    ws1 = wb.active
    ws1.title = "Performance Data"
    span = 6
    if sample:
        _banner(ws1, WATERMARK, span, _WM_FONT)
    _banner(ws1, "SIMULATED pre-deployment evidence — Financial Loss is a proxy, not incurred loss.", span, _NOTE_FONT)
    _write_header(
        ws1,
        [
            "Timestamp",
            "System ID",
            "Use Case ID",
            "Predicted Value",
            "Ground Truth Value",
            "Financial Loss Proxy (US$)",
        ],
    )
    for sys in systems:
        perf = sys.get("performance") or {}
        for run in perf.get("runs", []) or []:
            ws1.append(
                [
                    run.get("ts"),
                    sys.get("agent_id"),
                    run.get("scenario_id"),
                    run.get("actual"),        # model's predicted / produced value
                    run.get("expected"),      # ground truth
                    run.get("loss_proxy_usd", 0.0),
                ]
            )
    _autosize(ws1)

    # ---- Sheet 2: Covered Models ------------------------------------------
    ws2 = wb.create_sheet("Covered Models")
    span2 = 6
    if sample:
        _banner(ws2, WATERMARK, span2, _WM_FONT)
    _write_header(
        ws2,
        [
            "System ID",
            "Covered Model",
            "Function Scope",
            "Error Definition",
            "Unexpected-Error Threshold",
            "Proposed Error Limit (US$)",
        ],
    )
    for sys in systems:
        perf = sys.get("performance") or {}
        cm = perf.get("covered_model_draft")
        if cm:
            ws2.append(
                [
                    sys.get("agent_id"),
                    cm.get("covered_model"),
                    cm.get("function_scope"),
                    cm.get("error_definition"),
                    cm.get("unexpected_error_threshold"),
                    cm.get("proposed_error_limit_usd"),
                ]
            )
    _autosize(ws2)

    # ---- Sheet 3: Aggregates ----------------------------------------------
    ws3 = wb.create_sheet("Aggregates")
    span3 = 8
    if sample:
        _banner(ws3, WATERMARK, span3, _WM_FONT)
    _write_header(
        ws3,
        [
            "System ID",
            "Runs",
            "Error Rate",
            "CI 95% Low",
            "CI 95% High",
            "Total Loss Proxy (US$)",
            "Robustness Tests",
            "Robustness Pass Rate",
        ],
    )
    for sys in systems:
        perf = sys.get("performance") or {}
        agg = perf.get("aggregates") or {}
        ci = agg.get("ci_95", [None, None])
        rob = agg.get("robustness") or {}
        ws3.append(
            [
                sys.get("agent_id"),
                agg.get("n_runs"),
                agg.get("error_rate"),
                ci[0] if len(ci) > 0 else None,
                ci[1] if len(ci) > 1 else None,
                agg.get("total_loss_proxy_usd"),
                rob.get("n_tests"),
                rob.get("pass_rate"),
            ]
        )
    _autosize(ws3)

    out_path = Path(out_path)
    wb.save(out_path)
    return out_path
